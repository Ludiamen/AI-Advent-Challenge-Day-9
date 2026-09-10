"""Сводный отчёт по классификации в формате Excel.

Книга состоит из двух листов. «Классификация» — построчная расшифровка: что за
объект, куда его отнесли и почему. «Сводка» — итог по категориям с суммами,
количеством и долей.

Суммы на листе «Сводка» посчитаны в Python и записаны значениями. Формулы
SUMIF/COUNTIF со ссылками на соседний лист выглядят нарядно, но ломаются от
любого переименования листа или сдвига колонок, а проверить их без запуска
Excel нельзя. Значения же проверяются тестом.

Публичное API:
  build_pivot_report(items, output_path, cost_limit) -> str
"""

from __future__ import annotations

import os
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from agent.fsbu import CATEGORIES

SHEET_DETAILS = "Классификация"
SHEET_SUMMARY = "Сводка"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TOTAL_FONT = Font(bold=True)
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY = "#,##0.00"

# Заливка строки по категории — тот же смысловой код, что и в веб-интерфейсе.
_CATEGORY_FILL = {
    "НМА": PatternFill("solid", fgColor="E2EFDA"),            # зелёный
    "малоценный_НМА": PatternFill("solid", fgColor="FFF2CC"),  # жёлтый
    "расходы_периода": PatternFill("solid", fgColor="EDEDED"),  # серый
    "SaaS_подписка": PatternFill("solid", fgColor="DDEBF7"),    # синий
}

DETAIL_COLUMNS = [
    ("Объект учёта", 38),
    ("Категория", 18),
    ("Счёт", 10),
    ("Стоимость, ₽", 16),
    ("Срок, мес.", 11),
    ("СПИ, мес.", 11),
    ("Амортизация/мес., ₽", 20),
    ("Критерии (выполнено из 5)", 24),
    ("Рекомендация", 52),
    ("Риски", 42),
]

SUMMARY_COLUMNS = [("Категория", 22), ("Счёт", 12), ("Сумма, ₽", 16),
                   ("Количество", 13), ("Доля, %", 10)]


def _as_item_dict(item: Any) -> dict:
    """Принимает и объект Item, и обычный словарь — отчёт строится из обоих."""
    return item.to_dict() if hasattr(item, "to_dict") else dict(item)


def _write_header(worksheet, columns: list[tuple[str, int]]) -> None:
    for index, (title, width) in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=index, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"


def _fill_details(worksheet, items: list[dict]) -> None:
    _write_header(worksheet, DETAIL_COLUMNS)

    for row_number, item in enumerate(items, start=2):
        criteria = item.get("criteria") or {}
        met = sum(1 for value in criteria.values() if value)
        values = [
            item.get("position", ""),
            item.get("category", ""),
            item.get("account", ""),
            float(item.get("cost") or 0),
            item.get("term_months") or "",
            item.get("spi_months") or "",
            float(item.get("monthly_amortization") or 0),
            f"{met} из {len(criteria) or 5}",
            item.get("recommendation", ""),
            item.get("risk_notes", ""),
        ]
        fill = _CATEGORY_FILL.get(item.get("category", ""))
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=column in (1, 9, 10))
            if column in (4, 7):
                cell.number_format = _MONEY
            if fill is not None:
                cell.fill = fill

    total_row = len(items) + 2
    worksheet.cell(row=total_row, column=1, value="ИТОГО").font = _TOTAL_FONT
    total = worksheet.cell(row=total_row, column=4, value=sum(float(i.get("cost") or 0) for i in items))
    total.font = _TOTAL_FONT
    total.number_format = _MONEY


def _fill_summary(worksheet, items: list[dict], cost_limit: float) -> None:
    _write_header(worksheet, SUMMARY_COLUMNS)

    totals: dict[str, dict[str, float]] = {
        name: {"amount": 0.0, "count": 0} for name in CATEGORIES
    }
    for item in items:
        category = item.get("category", "")
        row = totals.setdefault(category, {"amount": 0.0, "count": 0})
        row["amount"] += float(item.get("cost") or 0)
        row["count"] += 1

    grand_total = sum(row["amount"] for row in totals.values())

    for row_number, (category, row) in enumerate(totals.items(), start=2):
        share = (row["amount"] / grand_total * 100) if grand_total else 0.0
        values = [
            category,
            CATEGORIES.get(category, {}).get("account", ""),
            row["amount"],
            int(row["count"]),
            round(share, 1),
        ]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.border = _BORDER
            if column == 3:
                cell.number_format = _MONEY
        fill = _CATEGORY_FILL.get(category)
        if fill is not None:
            worksheet.cell(row=row_number, column=1).fill = fill

    total_row = len(totals) + 2
    worksheet.cell(row=total_row, column=1, value="ИТОГО").font = _TOTAL_FONT
    total = worksheet.cell(row=total_row, column=3, value=grand_total)
    total.font = _TOTAL_FONT
    total.number_format = _MONEY
    worksheet.cell(row=total_row, column=4, value=len(items)).font = _TOTAL_FONT

    note_row = total_row + 2
    worksheet.cell(
        row=note_row,
        column=1,
        value=f"Лимит стоимости по учётной политике: {cost_limit:,.0f} ₽".replace(",", " "),
    )
    worksheet.cell(
        row=note_row + 1,
        column=1,
        value="Классификация по ФСБУ 14/2022 «Нематериальные активы» (приказ Минфина № 86н).",
    )


def build_pivot_report(items: list[Any], output_path: str, cost_limit: float = 100_000.0) -> str:
    """Собирает книгу Excel с расшифровкой и сводкой; возвращает путь к файлу."""
    rows = [_as_item_dict(item) for item in items]
    if not rows:
        raise ValueError("Нечего выгружать: список объектов пуст.")

    workbook = Workbook()
    _fill_details(workbook.active, rows)
    workbook.active.title = SHEET_DETAILS
    _fill_summary(workbook.create_sheet(SHEET_SUMMARY), rows, cost_limit)

    directory = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(directory, exist_ok=True)
    workbook.save(output_path)
    return output_path
