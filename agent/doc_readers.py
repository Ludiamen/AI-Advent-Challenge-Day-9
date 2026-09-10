"""Извлечение текста из документов: PDF, Word (.docx) и Excel (.xlsx).

Модуль отвечает ровно за одно: превратить файл в строку, пригодную для отправки
языковой модели. Он ничего не знает ни про ФСБУ, ни про LLM — и наоборот, агент
не знает, как устроен PDF. Это единственное место, где живут pypdf, python-docx
и openpyxl.

Публичное API:
  read_document(path, kind) -> str   — прочитать файл любого поддерживаемого типа
  detect_kind(path)         -> str   — определить тип по расширению
  SUPPORTED                          — поддерживаемые расширения
  DocumentError                      — единственный тип ошибки наружу
"""

from __future__ import annotations

import json
import os

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

# Расширение -> человекочитаемое название формата.
SUPPORTED: dict[str, str] = {
    ".pdf": "PDF",
    ".docx": "Word (.docx)",
    ".xlsx": "Excel (.xlsx)",
}

# Длинный документ обрезается: у моделей ограничен контекст, а платные тарифы
# считают каждый входной токен. Обрезка всегда помечается в тексте, чтобы это
# не выглядело так, будто документ кончился.
MAX_TEXT_LENGTH = 12_000
TRUNCATION_MARK = "\n\n[документ обрезан: показаны первые {n} символов]"

# Что вернуть, если PDF оказался сканом без текстового слоя.
NO_TEXT_LAYER = (
    "PDF не содержит текстового слоя (вероятно, скан). "
    "Требуется распознавание текста (OCR)."
)


class DocumentError(RuntimeError):
    """Понятная ошибка чтения документа для верхнего уровня."""


def detect_kind(path: str) -> str:
    """Определяет формат по расширению файла.

    Старый формат .doc не поддерживается намеренно: его чтение требует внешних
    конвертеров, а сообщение об этом полезнее молчаливой ошибки.
    """
    extension = os.path.splitext(path)[1].lower()
    if extension in SUPPORTED:
        return extension
    if extension == ".doc":
        raise DocumentError(
            "Старый формат .doc не поддерживается — пересохраните файл как .docx."
        )
    raise DocumentError(
        f"Неподдерживаемый формат «{extension or 'без расширения'}». "
        f"Допустимы: {', '.join(SUPPORTED)}."
    )


def truncate(text: str, limit: int = MAX_TEXT_LENGTH) -> str:
    """Обрезает длинный текст, оставляя явную пометку об этом."""
    if len(text) <= limit:
        return text
    return text[:limit] + TRUNCATION_MARK.format(n=limit)


def read_pdf(path: str, limit: int = MAX_TEXT_LENGTH) -> str:
    """Текст из PDF постранично; для скана возвращает объяснение вместо пустоты."""
    try:
        reader = PdfReader(path)
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:  # битый файл, шифрование, неверная структура
        raise DocumentError(f"Не удалось прочитать PDF: {exc}") from exc

    text = "\n".join(page for page in pages if page).strip()
    if not text:
        return NO_TEXT_LAYER
    return truncate(text, limit)


def read_docx(path: str, limit: int = MAX_TEXT_LENGTH) -> str:
    """Текст из Word: абзацы плюс содержимое таблиц.

    Таблицы читаются отдельно: в договорах и сметах именно там лежат позиции,
    суммы и сроки, а `doc.paragraphs` их не возвращает.
    """
    try:
        document = Document(path)
    except Exception as exc:
        raise DocumentError(f"Не удалось прочитать документ Word: {exc}") from exc

    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    for number, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"\nТаблица {number}:\n" + "\n".join(rows))

    return truncate("\n".join(parts).strip(), limit)


def read_xlsx(path: str, limit: int = MAX_TEXT_LENGTH) -> str:
    """Данные из Excel в виде JSON: строки как словари по заголовкам первой строки.

    JSON, а не «плоский» текст, потому что модели заметно проще связать значение
    с названием колонки, когда связь задана явно.
    """
    try:
        workbook = load_workbook(path, data_only=True)
    except Exception as exc:
        raise DocumentError(f"Не удалось прочитать книгу Excel: {exc}") from exc

    sheets: dict[str, list[dict]] = {}
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            continue

        headers = [
            str(cell).strip() if cell is not None else f"колонка_{i + 1}"
            for i, cell in enumerate(rows[0])
        ]
        records = []
        for row in rows[1:]:
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue  # пустая строка-разделитель
            records.append(
                {
                    header: (cell if cell is None else str(cell).strip())
                    for header, cell in zip(headers, row)
                }
            )
        if records:
            sheets[worksheet.title] = records

    if not sheets:
        return ""
    return truncate(json.dumps(sheets, ensure_ascii=False, indent=1), limit)


_READERS = {".pdf": read_pdf, ".docx": read_docx, ".xlsx": read_xlsx}


def read_document(path: str, kind: str = "", limit: int = MAX_TEXT_LENGTH) -> str:
    """Читает файл поддерживаемого формата и возвращает текст для модели.

    kind — расширение с точкой; если не задано, определяется по имени файла.
    Бросает DocumentError с понятным сообщением при любой проблеме.
    """
    if not os.path.exists(path):
        raise DocumentError(f"Файл не найден: {path}")

    kind = (kind or detect_kind(path)).lower()
    if kind not in _READERS:
        raise DocumentError(f"Неподдерживаемый формат «{kind}». Допустимы: {', '.join(SUPPORTED)}.")

    text = _READERS[kind](path, limit)
    if not text.strip():
        raise DocumentError(
            "Не удалось извлечь текст из документа: он пуст или не содержит данных."
        )
    return text
